"""Consultas agregadas y exportación a Excel / PDF.

Las mismas cifras alimentan la pantalla de reportes y los archivos que
descarga el staff: un solo lugar donde se define qué cuenta como
"renovación" o "recaudado".
"""

import io
from datetime import date
from decimal import Decimal

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.afiliado import Afiliado
from app.models.documento import Documento
from app.models.enums import CategoriaAfiliado, EstadoAfiliado, EstadoDocumento
from app.models.pago import Pago
from app.schemas.reporte import (
    BarraMes,
    FilaCategoria,
    FiltrosReporte,
    MetricaDashboard,
    ReporteResumen,
)

MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]

ETIQUETA_CATEGORIA = {
    CategoriaAfiliado.CLASE_A: "Clase A",
    CategoriaAfiliado.CLASE_B: "Clase B",
    CategoriaAfiliado.CLASE_C: "Clase C",
}

ETIQUETA_ESTADO = {
    EstadoAfiliado.ACTIVO: "Activo",
    EstadoAfiliado.PENDIENTE: "Pendiente",
    EstadoAfiliado.VENCIDO: "Vencido",
}


# --------------------------------------------------------------------------
# Métricas
# --------------------------------------------------------------------------


def metricas_dashboard(db: Session, dias_aviso: int = 30) -> MetricaDashboard:
    hoy = date.today()
    limite = date.fromordinal(hoy.toordinal() + dias_aviso)
    anio = hoy.year

    por_estado = dict(
        db.execute(select(Afiliado.estado, func.count()).group_by(Afiliado.estado)).all()
    )

    proximos = db.scalar(
        select(func.count())
        .select_from(Afiliado)
        .where(
            Afiliado.fecha_vencimiento.is_not(None),
            Afiliado.fecha_vencimiento >= hoy,
            Afiliado.fecha_vencimiento <= limite,
        )
    )
    por_revisar = db.scalar(
        select(func.count())
        .select_from(Documento)
        .where(Documento.estado == EstadoDocumento.PENDIENTE)
    )
    nuevos = db.scalar(
        select(func.count())
        .select_from(Afiliado)
        .where(func.extract("year", Afiliado.fecha_afiliacion) == anio)
    )
    recaudado = db.scalar(
        select(func.coalesce(func.sum(Pago.monto), 0)).where(
            func.extract("year", Pago.fecha) == anio
        )
    )
    recaudado_anterior = db.scalar(
        select(func.coalesce(func.sum(Pago.monto), 0)).where(
            func.extract("year", Pago.fecha) == anio - 1
        )
    )

    variacion = None
    if recaudado_anterior:
        variacion = float(
            (Decimal(recaudado) - Decimal(recaudado_anterior)) / Decimal(recaudado_anterior) * 100
        )

    return MetricaDashboard(
        total_afiliados=sum(por_estado.values()),
        activos=por_estado.get(EstadoAfiliado.ACTIVO, 0),
        pendientes=por_estado.get(EstadoAfiliado.PENDIENTE, 0),
        vencidos=por_estado.get(EstadoAfiliado.VENCIDO, 0),
        proximos_a_vencer=proximos or 0,
        documentos_por_revisar=por_revisar or 0,
        nuevos_en_el_ano=nuevos or 0,
        recaudado_periodo=Decimal(recaudado or 0),
        variacion_recaudacion=round(variacion, 1) if variacion is not None else None,
    )


def resumen(db: Session, filtros: FiltrosReporte) -> ReporteResumen:
    hoy = date.today()
    desde = filtros.desde or date(hoy.year, 1, 1)
    hasta = filtros.hasta or hoy

    cond_afiliado = [Afiliado.fecha_afiliacion.is_not(None)]
    if filtros.categoria:
        cond_afiliado.append(Afiliado.categoria == filtros.categoria)
    if filtros.estado:
        cond_afiliado.append(Afiliado.estado == filtros.estado)

    nuevas = db.scalar(
        select(func.count())
        .select_from(Afiliado)
        .where(*cond_afiliado, Afiliado.fecha_afiliacion.between(desde, hasta))
    )

    # Renovación = pago del período dentro del rango de un afiliado que ya
    # existía antes del rango.
    renovaciones = db.scalar(
        select(func.count(func.distinct(Pago.afiliado_id)))
        .join(Afiliado, Afiliado.id == Pago.afiliado_id)
        .where(
            Pago.fecha.between(desde, hasta),
            Afiliado.fecha_afiliacion < desde,
        )
    )
    bajas = db.scalar(
        select(func.count())
        .select_from(Afiliado)
        .where(
            Afiliado.estado == EstadoAfiliado.VENCIDO,
            Afiliado.fecha_vencimiento.between(desde, hasta),
        )
    )
    recaudado = db.scalar(
        select(func.coalesce(func.sum(Pago.monto), 0)).where(Pago.fecha.between(desde, hasta))
    )

    return ReporteResumen(
        tipo=filtros.tipo,
        desde=desde,
        hasta=hasta,
        nuevas_afiliaciones=nuevas or 0,
        renovaciones=renovaciones or 0,
        bajas=bajas or 0,
        recaudado=Decimal(recaudado or 0),
        por_mes=_barras_por_mes(db, desde, hasta),
        por_categoria=_filas_por_categoria(db, desde, hasta, filtros),
        total_afiliados=db.scalar(select(func.count()).select_from(Afiliado)) or 0,
    )


def _barras_por_mes(db: Session, desde: date, hasta: date) -> list[BarraMes]:
    filas = db.execute(
        select(
            func.extract("year", Afiliado.fecha_afiliacion).label("anio"),
            func.extract("month", Afiliado.fecha_afiliacion).label("mes"),
            Afiliado.estado,
            func.count(),
        )
        .where(Afiliado.fecha_afiliacion.between(desde, hasta))
        .group_by("anio", "mes", Afiliado.estado)
        .order_by("anio", "mes")
    ).all()

    # Año y mes van como dos columnas y no concatenados: SQLAlchemy numera los
    # parámetros del separador distinto en el SELECT y en el GROUP BY, y
    # Postgres deja de ver una sola expresión agrupada.
    recaudo = {
        (int(anio), int(mes)): total
        for anio, mes, total in db.execute(
            select(
                func.extract("year", Pago.fecha).label("anio"),
                func.extract("month", Pago.fecha).label("mes"),
                func.coalesce(func.sum(Pago.monto), 0),
            )
            .where(Pago.fecha.between(desde, hasta))
            .group_by("anio", "mes")
        ).all()
    }

    acumulado: dict[tuple[int, int], BarraMes] = {}
    for anio, mes, estado, total in filas:
        clave = (int(anio), int(mes))
        barra = acumulado.setdefault(
            clave, BarraMes(mes=MESES[int(mes) - 1], anio=int(anio))
        )
        if estado == EstadoAfiliado.ACTIVO:
            barra.activos = total
        elif estado == EstadoAfiliado.PENDIENTE:
            barra.pendientes = total
        else:
            barra.vencidos = total

    for clave, barra in acumulado.items():
        barra.recaudado = Decimal(recaudo.get(clave, 0) or 0)

    return [acumulado[c] for c in sorted(acumulado)]


def _filas_por_categoria(
    db: Session, desde: date, hasta: date, filtros: FiltrosReporte
) -> list[FilaCategoria]:
    conteos = db.execute(
        select(Afiliado.categoria, Afiliado.estado, func.count()).group_by(
            Afiliado.categoria, Afiliado.estado
        )
    ).all()
    recaudo = dict(
        db.execute(
            select(Afiliado.categoria, func.coalesce(func.sum(Pago.monto), 0))
            .join(Pago, Pago.afiliado_id == Afiliado.id)
            .where(Pago.fecha.between(desde, hasta))
            .group_by(Afiliado.categoria)
        ).all()
    )

    filas: dict[CategoriaAfiliado, FilaCategoria] = {
        cat: FilaCategoria(categoria=cat) for cat in CategoriaAfiliado
    }
    for categoria, estado, total in conteos:
        fila = filas[categoria]
        fila.total += total
        if estado == EstadoAfiliado.ACTIVO:
            fila.activos += total
        elif estado == EstadoAfiliado.PENDIENTE:
            fila.pendientes += total
        else:
            fila.vencidos += total

    for categoria, monto in recaudo.items():
        filas[categoria].recaudado = Decimal(monto or 0)

    orden = list(filas.values())
    if filtros.categoria:
        orden = [f for f in orden if f.categoria == filtros.categoria]
    return orden


# --------------------------------------------------------------------------
# Exportación
# --------------------------------------------------------------------------


def _afiliados_filtrados(db: Session, filtros: FiltrosReporte) -> list[Afiliado]:
    consulta = select(Afiliado).order_by(Afiliado.nombre)
    if filtros.estado:
        consulta = consulta.where(Afiliado.estado == filtros.estado)
    if filtros.categoria:
        consulta = consulta.where(Afiliado.categoria == filtros.categoria)
    return list(db.scalars(consulta).all())


def exportar_excel(db: Session, filtros: FiltrosReporte) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    afiliados = _afiliados_filtrados(db, filtros)
    datos = resumen(db, filtros)

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Afiliados"

    encabezados = [
        "Razón social",
        "RNC / Cédula",
        "Representante",
        "Categoría",
        "Estado",
        "Afiliación",
        "Vencimiento",
        "Cuota anual",
        "Correo",
        "Teléfono",
        "Contacto contabilidad",
        "Contacto marketing",
        "Contacto comercial",
    ]
    hoja.append(encabezados)

    relleno = PatternFill("solid", start_color="FF00776D")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A2"

    for a in afiliados:
        contactos = {c.area.value: c for c in a.contactos}

        def resumen_contacto(area: str) -> str:
            c = contactos.get(area)
            if not c:
                return "—"
            partes = [c.nombre]
            if c.cargo:
                partes.append(c.cargo)
            if c.telefono:
                partes.append(c.telefono)
            if c.email:
                partes.append(c.email)
            return " · ".join(partes)

        hoja.append(
            [
                a.nombre,
                a.rnc_cedula,
                a.representante or "—",
                ETIQUETA_CATEGORIA[a.categoria],
                ETIQUETA_ESTADO[a.estado],
                a.fecha_afiliacion,
                a.fecha_vencimiento,
                float(a.cuota_anual) if a.cuota_anual is not None else None,
                a.email or "—",
                a.telefono or "—",
                resumen_contacto("contabilidad"),
                resumen_contacto("marketing"),
                resumen_contacto("comercial"),
            ]
        )

    for col, ancho in enumerate(
        [34, 16, 22, 12, 12, 13, 13, 14, 26, 16, 40, 40, 40], start=1
    ):
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    for fila in hoja.iter_rows(min_row=2, min_col=6, max_col=7):
        for celda in fila:
            celda.number_format = "DD/MM/YYYY"
    for fila in hoja.iter_rows(min_row=2, min_col=8, max_col=8):
        for celda in fila:
            celda.number_format = '"RD$" #,##0.00'

    hoja2 = wb.create_sheet("Resumen")
    hoja2.append(["Reporte", datos.tipo.value])
    hoja2.append(["Desde", datos.desde])
    hoja2.append(["Hasta", datos.hasta])
    hoja2.append([])
    hoja2.append(["Nuevas afiliaciones", datos.nuevas_afiliaciones])
    hoja2.append(["Renovaciones", datos.renovaciones])
    hoja2.append(["Bajas", datos.bajas])
    hoja2.append(["Recaudado", float(datos.recaudado)])
    hoja2.append([])
    hoja2.append(["Categoría", "Afiliados", "Activos", "Pendientes", "Vencidos", "Recaudado"])
    for fila in datos.por_categoria:
        hoja2.append(
            [
                ETIQUETA_CATEGORIA[fila.categoria],
                fila.total,
                fila.activos,
                fila.pendientes,
                fila.vencidos,
                float(fila.recaudado),
            ]
        )
    hoja2.column_dimensions["A"].width = 22
    for celda in hoja2[10]:
        celda.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def exportar_pdf(db: Session, filtros: FiltrosReporte) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    afiliados = _afiliados_filtrados(db, filtros)
    datos = resumen(db, filtros)

    tinta = colors.HexColor("#233738")
    teal = colors.HexColor("#00776d")
    hueso = colors.HexColor("#fcfcf7")
    gris = colors.HexColor("#e9e9e9")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(LETTER),
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"ADECLA · Reporte de {datos.tipo.value}",
    )
    base = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "TituloAdecla", parent=base["Title"], textColor=tinta, fontSize=18, alignment=0
    )
    sub = ParagraphStyle(
        "SubAdecla", parent=base["Normal"], textColor=colors.HexColor("#5b6b6c"), fontSize=9
    )

    elementos = [
        Paragraph("ADECLA · Reporte de afiliados", titulo),
        Paragraph(
            f"{datos.tipo.value.capitalize()} · "
            f"{datos.desde:%d/%m/%Y} — {datos.hasta:%d/%m/%Y} · "
            f"{len(afiliados)} registros",
            sub,
        ),
        Spacer(1, 8 * mm),
    ]

    resumen_tabla = Table(
        [
            ["NUEVAS AFILIACIONES", "RENOVACIONES", "BAJAS", "RECAUDADO"],
            [
                str(datos.nuevas_afiliaciones),
                str(datos.renovaciones),
                str(datos.bajas),
                f"RD$ {datos.recaudado:,.2f}",
            ],
        ],
        colWidths=[60 * mm] * 4,
    )
    resumen_tabla.setStyle(
        TableStyle(
            [
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5b6b6c")),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 15),
                ("TEXTCOLOR", (0, 1), (-1, 1), tinta),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
                ("LINEBELOW", (0, 1), (-1, 1), 0.5, gris),
            ]
        )
    )
    elementos += [resumen_tabla, Spacer(1, 8 * mm)]

    filas = [["RAZÓN SOCIAL", "RNC", "REPRESENTANTE", "CATEGORÍA", "VENCE", "ESTADO"]]
    for a in afiliados:
        filas.append(
            [
                a.nombre,
                a.rnc_cedula,
                a.representante or "—",
                ETIQUETA_CATEGORIA[a.categoria],
                a.fecha_vencimiento.strftime("%d/%m/%y") if a.fecha_vencimiento else "—",
                ETIQUETA_ESTADO[a.estado],
            ]
        )

    tabla = Table(filas, colWidths=[75 * mm, 32 * mm, 48 * mm, 28 * mm, 24 * mm, 26 * mm], repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TEXTCOLOR", (0, 1), (-1, -1), tinta),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hueso]),
                ("GRID", (0, 0), (-1, -1), 0.25, gris),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    return buffer
